"""E-Reporting: angka laporan sesuai data, Excel valid, cetak, simpan ke arsip; lingkup mengikuti hak akses."""
import io
from datetime import timedelta

from django.utils import timezone

from arsip.models import ArchiveDocument
from workflow import services as alur

from . import skenario
from ._bantu import JUDUL_QA, berkas, url


@skenario("laporan", "Laporan per periode/unit: angka benar, Excel (.xlsx), cetak, simpan ke E-Arsip", crud=True,
          role="ka_igd", menu=["laporan:index", "laporan:xlsx", "laporan:archive"])
def _(u):
    w, today = u.w, timezone.localdate()
    buat = lambda judul, due: alur.create_task(created_by=w.ka_igd, source_type="LANGSUNG", source_note="QA",  # noqa: E731
                                               title=f"{JUDUL_QA} {judul}", unit=w.igd, pic=w.staf_igd,
                                               due_date=today + timedelta(days=due))
    selesai = buat("Lap selesai", 3)
    alur.accept_task(selesai, w.staf_igd)
    alur.report_progress(selesai, w.staf_igd, progress=100, status="BERJALAN")
    alur.add_evidence(selesai, w.staf_igd, berkas())
    alur.submit_result(selesai, w.staf_igd)
    alur.verify(selesai, w.ka_igd, approved=True)
    buat("Lap terlambat", -2)
    buat("Lap berjalan", 5)
    q = f"?dari={today.isoformat()}&sampai={today.isoformat()}&unit={w.igd.pk}"
    r = u.get(url("laporan:index") + q, ada=["Evaluasi per unit", "Dibuat", "Total tugas"])
    from laporan.services import build

    rep = build(w.ka_igd, dari=today, sampai=today, unit=w.igd)
    rk = rep["ringkas"]
    u.harap((rk["total"], rk["selesai"], rk["terlambat_terbuka"], rk["pct_tepat_waktu"]) == (3, 1, 1, 100),
            f"ringkasan salah: {dict(rk)}")
    u.harap(rk["selesai_terverifikasi"] == 1, "selesai terverifikasi tidak dibedakan")
    x = u.get(url("laporan:xlsx") + q, harap=200)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(b"".join(x.streaming_content) if x.streaming else x.content))
    u.harap(wb.sheetnames == ["Ringkasan", "Per unit", "Per jenis", "Per PIC"], f"lembar: {wb.sheetnames}")
    nilai = {row[0].value: row[1].value for row in wb["Ringkasan"].iter_rows() if row[0].value}
    u.harap(nilai.get("Total tugas") == 3, f"Excel total {nilai.get('Total tugas')}")
    u.get(url("laporan:index") + q + "&cetak=1", ada=["LAPORAN PELAKSANAAN", "paper"])
    u.post(url("laporan:archive"), {"dari": today.isoformat(), "sampai": today.isoformat(), "unit": w.igd.pk},
           harap=302)
    doc = ArchiveDocument.objects.filter(doc_type="LAPORAN", owner=w.ka_igd).get()
    u.get(url("arsip:detail", doc.pk), ada=[".xlsx"])
    u.get(url("laporan:index") + "?dari=2026-12-31&sampai=2026-01-01", ada=["Tanggal awal harus sebelum"])
    u.sebagai("staf_farmasi")                          # lingkup lain: tidak melihat tugas IGD
    rep2 = build(w.staf_farmasi, dari=today, sampai=today)
    u.harap(rep2["ringkas"]["total"] == 0, "staf farmasi melihat tugas IGD di laporan")
    del r
