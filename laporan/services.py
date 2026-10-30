"""E-Reporting: laporan pelaksanaan & evaluasi unit dari data sistem (PRD 6.11).

Periode = tugas yang DIBUAT dalam rentang tanggal. Semua angka mengikuti lingkup akses pengguna
(visible_tasks / visible_instructions / visible_dispositions) — sama dengan dashboard."""
import io
from collections import OrderedDict

from django.db.models import Count, Q
from django.utils import timezone

from instructions.models import Instruction
from instructions.services import visible_instructions
from letters.models import Disposition
from letters.services import visible_dispositions
from workflow.access import visible_tasks
from workflow.models import CLOSED_STATUSES, TaskSource, TaskStatus

S = TaskStatus


def _ringkas_tugas(qs, today):
    data = list(qs.values("id", "status", "due_date", "created_at", "closed_at", "submitted_at",
                          "requires_verification"))
    selesai = [t for t in data if t["status"] == S.SELESAI]
    tepat = [t for t in selesai if t["due_date"] and timezone.localtime(t["submitted_at"] or t["closed_at"]).date()
             <= t["due_date"]]
    berbatas = [t for t in selesai if t["due_date"]]
    durasi = [((t["closed_at"] - t["created_at"]).total_seconds() / 86400) for t in selesai if t["closed_at"]]
    terbuka = [t for t in data if t["status"] not in CLOSED_STATUSES]
    return OrderedDict([
        ("total", len(data)),
        ("terbuka", len(terbuka)),
        ("selesai", len(selesai)),
        ("selesai_terverifikasi", sum(1 for t in selesai if t["requires_verification"])),
        ("selesai_tanpa_verifikasi", sum(1 for t in selesai if not t["requires_verification"])),
        ("menunggu_verifikasi", sum(1 for t in data if t["status"] == S.MENUNGGU_VERIFIKASI)),
        ("perlu_perbaikan", sum(1 for t in data if t["status"] == S.PERLU_PERBAIKAN)),
        ("terhambat", sum(1 for t in data if t["status"] == S.TERHAMBAT)),
        ("terlambat_terbuka", sum(1 for t in terbuka if t["due_date"] and t["due_date"] < today)),
        ("selesai_terlambat", len(berbatas) - len(tepat)),
        ("dibatalkan", sum(1 for t in data if t["status"] == S.DIBATALKAN)),
        ("pct_tepat_waktu", round(100 * len(tepat) / len(berbatas)) if berbatas else None),
        ("pct_selesai", round(100 * len(selesai) / (len(data) - sum(1 for t in data if t["status"] == S.DIBATALKAN)))
         if len(data) - sum(1 for t in data if t["status"] == S.DIBATALKAN) else None),
        ("rata_hari", round(sum(durasi) / len(durasi), 1) if durasi else None),
    ])


LABEL = OrderedDict([
    ("total", "Total tugas"), ("terbuka", "Belum ditutup"), ("selesai", "Selesai (ditutup)"),
    ("selesai_terverifikasi", "— selesai & diverifikasi"), ("selesai_tanpa_verifikasi", "— selesai tanpa verifikasi"),
    ("menunggu_verifikasi", "Menunggu verifikasi (belum ditutup)"), ("perlu_perbaikan", "Perlu perbaikan"),
    ("terhambat", "Terhambat"), ("terlambat_terbuka", "Terlambat (masih terbuka)"),
    ("selesai_terlambat", "Selesai melewati tenggat"), ("dibatalkan", "Dibatalkan"),
    ("pct_tepat_waktu", "Selesai tepat waktu (%)"), ("pct_selesai", "Tingkat penyelesaian (%)"),
    ("rata_hari", "Rata-rata waktu penyelesaian (hari)"),
])


def build(user, *, dari, sampai, unit=None, pic=None, sumber="", status="", today=None):
    today = today or timezone.localdate()
    qs = visible_tasks(user).filter(created_at__date__gte=dari, created_at__date__lte=sampai)
    if unit:
        qs = qs.filter(unit=unit)
    if pic:
        qs = qs.filter(pic=pic)
    if sumber:
        qs = qs.filter(source_type=sumber)
    if status:
        qs = qs.filter(status=status)
    qs = qs.distinct()
    per_unit = []
    for u in qs.values("unit_id", "unit__code", "unit__name").distinct().order_by("unit__code"):
        r = _ringkas_tugas(qs.filter(unit_id=u["unit_id"]), today)
        per_unit.append({"kode": u["unit__code"], "nama": u["unit__name"], **r})
    per_sumber = []
    for kode, nama in TaskSource.choices:
        sub = qs.filter(source_type=kode)
        if sub.exists():
            per_sumber.append({"nama": nama, **_ringkas_tugas(sub, today)})
    per_pic = []
    for p in qs.values("pic_id", "pic__first_name", "pic__last_name", "pic__username").distinct():
        nama = f"{p['pic__first_name']} {p['pic__last_name']}".strip() or p["pic__username"]
        per_pic.append({"nama": nama, **_ringkas_tugas(qs.filter(pic_id=p["pic_id"]), today)})
    per_pic.sort(key=lambda r: (-(r["terlambat_terbuka"]), r["nama"]))
    ins = visible_instructions(user).filter(instruction_date__gte=dari, instruction_date__lte=sampai)
    disp = visible_dispositions(user).filter(created_at__date__gte=dari, created_at__date__lte=sampai)
    return {
        "ringkas": _ringkas_tugas(qs, today), "label": LABEL,
        "per_unit": per_unit, "per_sumber": per_sumber, "per_pic": per_pic,
        "instruksi": dict(ins.values_list("status").annotate(n=Count("id"))),
        "instruksi_label": dict(Instruction.Status.choices),
        "disposisi": dict(disp.values_list("status").annotate(n=Count("id"))),
        "disposisi_label": dict(Disposition.Status.choices),
        "tindak_lanjut": qs.filter(source_type=TaskSource.KEPUTUSAN_RAPAT).aggregate(
            total=Count("id"), selesai=Count("id", filter=Q(status=S.SELESAI))),
        "dibuat": timezone.localtime(),
        "parameter": OrderedDict([("Periode", f"{dari:%d-%m-%Y} s.d. {sampai:%d-%m-%Y}"),
                                  ("Unit", str(unit) if unit else "Semua"), ("PIC", str(pic) if pic else "Semua"),
                                  ("Jenis pekerjaan", dict(TaskSource.choices).get(sumber, "Semua")),
                                  ("Status", dict(TaskStatus.choices).get(status, "Semua")),
                                  ("Lingkup data", f"sesuai hak akses {user} ({user.get_role_display()})")]),
        "tugas_terlambat": qs.overdue(today).select_related("pic", "unit").order_by("due_date")[:50],
    }


def xlsx(rep):
    """Laporan -> berkas Excel (.xlsx) dengan parameter & waktu pembuatan di setiap lembar."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    tebal, latar = Font(bold=True), PatternFill("solid", fgColor="E1EFFC")

    def kepala(ws, judul):
        ws.append([judul])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([f"Dibuat {rep['dibuat']:%d-%m-%Y %H:%M}"])
        for k, v in rep["parameter"].items():
            ws.append([k, v])
        ws.append([])

    ws = wb.active
    ws.title = "Ringkasan"
    kepala(ws, "Laporan Pelaksanaan Pekerjaan — E-Office")
    for k, lbl in rep["label"].items():
        ws.append([lbl, rep["ringkas"][k]])
    ws.append([])
    ws.append(["Instruksi per status"])
    for k, n in rep["instruksi"].items():
        ws.append([rep["instruksi_label"].get(k, k), n])
    ws.append(["Disposisi per status"])
    for k, n in rep["disposisi"].items():
        ws.append([rep["disposisi_label"].get(k, k), n])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 60

    kolom = list(rep["label"].keys())
    for judul, baris, kunci in (("Per unit", rep["per_unit"], "kode"), ("Per jenis", rep["per_sumber"], "nama"),
                                ("Per PIC", rep["per_pic"], "nama")):
        w = wb.create_sheet(judul)
        kepala(w, f"Laporan {judul.lower()}")
        w.append(["Unit" if kunci == "kode" else judul.split()[1].capitalize()] + [rep["label"][k] for k in kolom])
        for c in w[w.max_row]:
            c.font, c.fill = tebal, latar
        for r in baris:
            w.append([r[kunci]] + [r[k] for k in kolom])
        w.column_dimensions["A"].width = 28
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
